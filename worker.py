"""Offline typed Laya decisions and read-only Excel import. No per-module recipes."""

import os, sys, json, time, contextlib, re

os.environ.setdefault('USE_TF', '0')
os.environ.setdefault('TOKENIZERS_PARALLELISM', 'false')
for stream in (sys.stdin, sys.stdout, sys.stderr):
    if hasattr(stream, 'reconfigure'):
        stream.reconfigure(encoding='utf-8')
ALIASES = {
    'id': ['测试编号', '用例编号', '编号', 'caseid', 'id'],
    'title': ['测试用例名称', '测试用例', '用例名称', '标题', 'title', 'name', 'testcase'],
    'steps': ['测试步骤', '操作步骤', '步骤', 'steps', 'teststeps'],
    'expected': ['预期结果', '期望结果', '预期', 'expected', 'expectedresult'],
    'preconditions': ['前提条件', '前置条件', '前提', 'preconditions'],
    'data': ['测试数据', '数据', 'testdata', 'data'],
    'url': ['页面地址', '网址', 'url', 'pageurl'],
    'source_result': [
        '实际执行结果',
        '实际结果',
        '测试结果',
        'actualresult',
        'testresult',
        'result',
    ],
}


def normalize(s):
    return re.sub(r'[\s_\-]', '', str(s or '')).lower()


def read_excel(filename):
    from openpyxl import load_workbook

    wb = load_workbook(filename, read_only=True, data_only=False)
    cases = []
    warnings = []
    for sheet in wb:
        mapping = None
        for rowno, row in enumerate(sheet.iter_rows(), 1):
            vals = [str(c.value if c.value is not None else '') for c in row]
            if mapping is None:
                names = [normalize(v) for v in vals]
                found = {
                    key: next((i for i, n in enumerate(names) if n in synonyms), None)
                    for key, synonyms in ALIASES.items()
                }
                if found['steps'] is not None and found['expected'] is not None:
                    mapping = found
                    continue
                if rowno >= 15:
                    break
                continue
            get = lambda key: (
                vals[mapping[key]]
                if mapping.get(key) is not None and mapping[key] < len(vals)
                else ''
            )
            if not any(v.strip() for v in vals):
                continue
            if not get('steps') and not get('expected'):
                continue
            cid = get('id') or f'{sheet.title}-{rowno}'
            if cid.isdigit():
                cid = cid.zfill(3)
            formulas = [
                row[i].coordinate
                for i in mapping.values()
                if i is not None and i < len(row) and row[i].data_type == 'f'
            ]
            item = {
                **{key: get(key) for key in ALIASES},
                'id': cid,
                'sheet': sheet.title,
                'row': rowno,
                'key': f'{sheet.title}:{rowno}',
                'import_error': ('执行列含公式，需提供明确文本：' + ','.join(formulas))
                if formulas
                else '',
            }
            # Common merged-ID layout: blank ID/title rows continue the preceding case.
            if not get('id') and not get('title') and cases and cases[-1]['sheet'] == sheet.title:
                for key in ('steps', 'expected', 'preconditions', 'data', 'import_error'):
                    if item[key]:
                        cases[-1][key] += ('\n' if cases[-1][key] else '') + item[key]
            else:
                cases.append(item)
        if mapping is None:
            warnings.append(f'{sheet.title}:前15行找不到步骤与预期结果表头，未读取')
    wb.close()
    return {'cases': cases, 'warnings': warnings}


agent = None
for line in sys.stdin:
    req = {}
    try:
        req = json.loads(line)
        if req['action'] == 'read_excel':
            result = read_excel(req['path'])
        else:
            if agent is None:
                start = time.perf_counter()
                with contextlib.redirect_stdout(sys.stderr):
                    import laya, torch

                    torch.set_num_threads(4)
                    agent = laya.load(req['model'], device='cpu')
                load_ms = round((time.perf_counter() - start) * 1000, 2)
            if req['action'] == 'load':
                result = {'load_ms': load_ms, 'device': 'cpu'}
            elif req['action'] == 'decide':
                if not 2 <= len(req['criteria']) <= 8:
                    raise ValueError('每次选择必须有2至8个候选，包括无法判断选项')
                start = time.perf_counter()
                with contextlib.redirect_stdout(sys.stderr):
                    answer = agent.predict(
                        req['state'],
                        {
                            'scenario': {
                                'type': 'choice',
                                'instructions': '用户请求对应哪个操作？',
                                'criteria': req['criteria'],
                            }
                        },
                    )['answers']['scenario']
                result = {**answer, 'inference_ms': round((time.perf_counter() - start) * 1000, 2)}
            else:
                raise ValueError('不支持的请求')
        print(json.dumps({'id': req['id'], **result}, ensure_ascii=False), flush=True)
    except Exception as e:
        print(
            json.dumps(
                {'id': req.get('id'), 'error': f'{type(e).__name__}: {e}'}, ensure_ascii=False
            ),
            flush=True,
        )
