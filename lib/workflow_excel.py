"""Standalone XLSX contract for generated browser workflows."""

import copy
import hashlib
import json
import math
import sys
import unicodedata
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter


HEADERS = [
    '测试编号',
    '模块名称',
    '子模块名称',
    '测试用例名称',
    '优先级',
    '用例类型',
    '用例类别',
    '前提条件',
    '测试步骤',
    '预期结果',
    '实际执行结果',
    '执行测试人员',
    '执行状态',
    '用例引入阶段',
    '执行日期',
    '备注',
    '是否提bug',
]
META = '__laya_steps__'
TITLES = {
    'tabs': '页签入口可见',
    'tab-switch': '页签切换后可见并可切回',
    'table': '列表列头可见',
    'filters': '名称筛选控件可见',
    'form-validation': '新增表单必填项为空时阻止保存',
    'form-invalid': '数值区间示例的非法括号触发校验',
    'create': '新增独立测试记录',
    'search': '按名称查询刚创建的记录',
    'view': '查看独立测试记录',
    'edit': '修改独立测试记录名称',
    'delete': '删除独立测试记录',
}
OPERATIONS = {'create': '新增', 'search': '查询', 'view': '查看', 'edit': '修改', 'delete': '删除'}


def cell_text(value):
    return '' if value is None else clean_text(str(value))


def clean_text(value):
    if not isinstance(value, str):
        return value
    value = ILLEGAL_CHARACTERS_RE.sub('', value)
    # JavaScript UTF-16 slicing can leave a lone surrogate at a truncation boundary.
    # Replacing it before hashing or writing keeps XLSX/XML encoding valid.
    return ''.join('\ufffd' if 0xD800 <= ord(char) <= 0xDFFF else char for char in value)


def visible_fields(sheet, row):
    return [cell_text(sheet.cell(row, col).value) for col in (1, 4, 8, 9, 10)]


def digest(fields):
    return hashlib.sha256(
        json.dumps(fields, ensure_ascii=False, separators=(',', ':')).encode()
    ).hexdigest()


def step_text(step):
    kind = step['kind']
    target = step.get('target', {}).get('name', '')
    value = step.get('value', '')
    if kind == 'click':
        return f'点击「{target}」' + (
            '（本轮记录所在行）' if step.get('target', {}).get('scope') == 'owned-row' else ''
        )
    if kind == 'fill':
        if isinstance(value, dict):
            value = '${' + value['fixture'] + '}'
        return f'在「{target}」输入「{value}」'
    if kind == 'form-fill':
        return f'在表单「{step["target"]["label"]}」第{step["target"]["index"] + 1}个输入框输入「{value}」'
    if kind == 'form-fill-required':
        return f'在必填输入框「{step["target"]["placeholder"]}」第{step["target"]["index"] + 1}项输入「{value}」'
    if kind == 'form-select':
        return f'在表单「{step["target"]["label"]}」依次选择：' + ' → '.join(step['path'])
    if kind == 'assert-form-errors':
        return '核对表单未关闭，必填字段显示校验错误：' + '、'.join(value)
    if kind == 'assert-input-invalid':
        return f'核对必填输入框「{step["target"]["placeholder"]}」第{step["target"]["index"] + 1}项显示格式错误'
    if kind == 'reload':
        return '重新加载页面，验证数据已保存'
    if kind == 'assert-row':
        return f'核对列表存在「{value}」'
    if kind == 'assert-absent':
        return f'核对列表不存在「{value}」'
    if kind == 'assert-text':
        return f'核对详情显示「{value}」'
    if kind == 'assert-control':
        return f'核对「{target}」控件可见'
    if kind == 'assert-selected':
        return f'核对「{target}」Tab已选中'
    if kind == 'assert-headers':
        return '核对列表列头包含：' + '、'.join(value)
    if kind == 'assert-disabled':
        return f'核对必填项为空时「{target}」按钮处于禁用状态（阻止保存）'
    raise ValueError('不支持的步骤：' + kind)


def expected_text(case):
    checks = [step_text(step) for step in case['steps'] if step['kind'].startswith('assert-')]
    if not checks:
        raise ValueError('用例没有页面断言：' + case['id'])
    return '\n'.join(f'{i}、{text}' for i, text in enumerate(checks, 1))


def load_template(path):
    if not path:
        source = Workbook()
        sheet = source.active
        sheet.title = '测试用例'
        sheet.append(HEADERS)
        return source, sheet
    path = Path(path)
    if not path.is_file():
        raise ValueError('找不到原 Excel 模板，请传 --template-excel：' + str(path))
    source = load_workbook(path, read_only=False)
    sheet = source.active
    actual = [sheet.cell(1, i).value for i in range(1, 18)]
    if actual != HEADERS:
        raise ValueError('模板表头不是之前的 17 列用例格式')
    return source, sheet


def write(payload):
    file = payload['file']

    # 失败原因常含 ANSI 转义序列（\u001b[2m 等）与其它控制字符，openpyxl 会抛
    # IllegalCharacterError 并把整个用例文件写出炸掉。统一在入口清洗。
    def clean(v):
        return clean_text(v)

    source, old = load_template(payload.get('template'))
    book = Workbook()
    sheet = book.active
    sheet.title = '生成用例'
    sheet.sheet_view.showGridLines = old.sheet_view.showGridLines
    sheet.freeze_panes = 'A2'
    for col in range(1, 18):
        original = old.cell(1, col)
        cell = sheet.cell(1, col, original.value)
        if original.has_style:
            cell._style = copy.copy(original._style)
        cell.alignment = copy.copy(original.alignment)
        letter = get_column_letter(col)
        sheet.column_dimensions[letter].width = old.column_dimensions[letter].width or 16
        # The source hides several case-identifying columns. Generated cases
        # should expose their ID and module while retaining source widths.
        sheet.column_dimensions[letter].hidden = False
    sheet.row_dimensions[1].height = old.row_dimensions[1].height
    meta = book.create_sheet(META)
    meta.sheet_state = 'veryHidden'
    meta['A1'] = clean_text(
        json.dumps(
            {
                'schemaVersion': file['schemaVersion'],
                'generatedAt': file['generatedAt'],
                'moduleUrl': file['moduleUrl'],
                'coverage': file['coverage'],
            },
            ensure_ascii=False,
        )
    )
    meta.append(['id', 'visible_sha256', 'case_json'])
    module = Path(file['moduleUrl'].split('?')[0].rstrip('/')).name or '目标页面'
    for index, case in enumerate(file['cases'], 2):
        uses_owned_record = case['operation'] in ('create', 'search', 'view', 'edit', 'delete')
        precondition = (
            '已登录并进入目标模块；使用本轮唯一命名的独立测试记录'
            if uses_owned_record
            else '已登录并进入目标模块'
        )
        remark = (
            '生成时已执行并验证；执行模式将使用新的独立测试记录'
            if uses_owned_record
            else '生成时已执行并验证；执行模式将重新检查页面，不创建记录'
        )
        values = [
            f'{index - 1:03d}',
            module,
            module,
            f'{module} - {TITLES[case["operation"]]}',
            '中',
            '正案例',
            '功能',
            precondition,
            '\n'.join(f'{n}、{step_text(step)}' for n, step in enumerate(case['steps'], 1)),
            expected_text(case),
            None,
            None,
            '未执行',
            '自动生成',
            None,
            remark,
            None,
        ]
        values = [clean(v) for v in values]
        for col, value in enumerate(values, 1):
            dest = sheet.cell(index, col, value)
            sample = old.cell(2, col)
            if sample.has_style:
                dest._style = copy.copy(sample._style)
            dest.alignment = copy.copy(sample.alignment)
            if col == 1:
                dest.number_format = '@'
            elif isinstance(value, str):
                # Names and page text are untrusted workbook input. Keep them
                # literal even if a site label begins with an Excel formula.
                dest.data_type = 's'
        sheet.row_dimensions[index].height = min(
            280, max(42, 18 * max(values[8].count('\n') + 1, values[9].count('\n') + 1))
        )
        meta.append(
            [
                values[0],
                digest(visible_fields(sheet, index)),
                clean(json.dumps(case, ensure_ascii=False)),
            ]
        )
    sheet.auto_filter.ref = f'A1:Q{max(2, sheet.max_row)}'
    Path(payload['path']).parent.mkdir(parents=True, exist_ok=True)
    book.save(payload['path'])
    source.close()
    return {'path': payload['path'], 'cases': len(file['cases'])}


def read(payload):
    book = load_workbook(payload['path'], read_only=False, data_only=False)
    if META not in book or book[META].sheet_state not in ('hidden', 'veryHidden'):
        raise ValueError('没有隐藏的 Laya 回放步骤，不能仅凭可见文字执行写入操作')
    sheet = book['生成用例']
    if [sheet.cell(1, i).value for i in range(1, 18)] != HEADERS:
        raise ValueError('生成用例表头已改变，无法安全读取')
    meta = book[META]
    header = json.loads(meta['A1'].value)
    cases = []
    for row in range(3, meta.max_row + 1):
        case_id, saved_hash, recipe = [meta.cell(row, col).value for col in (1, 2, 3)]
        visible_row = row - 1
        if digest(visible_fields(sheet, visible_row)) != saved_hash:
            raise ValueError(
                f'第{visible_row}行可见用例已修改，与隐藏回放步骤不一致，请重新生成或校对后再执行'
            )
        case = json.loads(recipe)
        if case_id != sheet.cell(visible_row, 1).value:
            raise ValueError(f'第{visible_row}行用例编号不一致')
        cases.append(case)
    book.close()
    return {
        'schemaVersion': header['schemaVersion'],
        'moduleUrl': header['moduleUrl'],
        'cases': cases,
        'coverage': header.get('coverage', []),
    }


def read_editable(payload):
    book = load_workbook(payload['path'], read_only=False, data_only=False)
    if META not in book or book[META].sheet_state not in ('hidden', 'veryHidden'):
        raise ValueError('没有隐藏的 Laya 回放步骤，不能执行重新绑定')
    sheet = book['生成用例']
    if [sheet.cell(1, i).value for i in range(1, 18)] != HEADERS:
        raise ValueError('生成用例表头已改变，无法重新绑定')
    meta = book[META]
    header = json.loads(meta['A1'].value)
    rows = []
    for row in range(3, meta.max_row + 1):
        case_id, saved_hash, recipe = [meta.cell(row, col).value for col in (1, 2, 3)]
        visible_row = row - 1
        fields = visible_fields(sheet, visible_row)
        rows.append(
            {
                'id': case_id,
                'visible': fields,
                'changed': digest(fields) != saved_hash,
                'case': json.loads(recipe),
            }
        )
    book.close()
    return {
        'schemaVersion': header['schemaVersion'],
        'moduleUrl': header['moduleUrl'],
        'coverage': header.get('coverage', []),
        'rows': rows,
    }


def results(payload):
    book = load_workbook(payload['path'])
    sheet = book['生成用例']
    sheet.column_dimensions['P'].width = max(sheet.column_dimensions['P'].width or 0, 70)
    for row, result in enumerate(payload['results'], 2):
        if (
            row > sheet.max_row
            or result['id'] != json.loads(book[META].cell(row + 1, 3).value)['id']
        ):
            raise ValueError('执行结果与生成用例顺序不一致')
        sheet.cell(row, 11).value = (
            'pass'
            if result['status'] == '通过'
            else 'skip'
            if result['status'] == '未执行（依赖阻断）'
            else 'fail'
        )
        sheet.cell(row, 13).value = (
            '未执行' if result['status'] == '未执行（依赖阻断）' else '已执行'
        )
        sheet.cell(row, 15).value = (
            None if result['status'] == '未执行（依赖阻断）' else datetime.now()
        )
        diagnosis = result.get('diagnosis') or {}
        step = result.get('failedStep') or {}
        notes = [f'原因分类：{diagnosis["category"]}'] if diagnosis.get('category') else []
        if step.get('number'):
            notes.append(
                f'失败步骤：{step["number"]}（{step.get("action", "")} {step.get("target", "")}）'
            )
        if not diagnosis.get('primary') or diagnosis['primary'] not in result['reason']:
            notes.append(f'执行错误：{result["reason"]}')
        if diagnosis.get('assessment'):
            notes.append(f'判断：{diagnosis["assessment"]}')
        if diagnosis.get('primary'):
            notes.append(f'同期证据：{diagnosis["primary"]}')
        for attempt in result.get('attempts', []):
            notes.append(
                f'尝试：{attempt["step"]} 第{attempt["attempt"]}次 {attempt["status"]} {attempt.get("reason", "")}'
            )
        note = '\n'.join(notes)[:3000]
        # 回放失败原因含 ANSI 转义（\u001b[2m 等），openpyxl 会抛 IllegalCharacterError
        note = clean_text(note)
        remark = sheet.cell(row, 16)
        remark.value = note
        alignment = copy.copy(remark.alignment)
        alignment.wrap_text = True
        alignment.vertical = 'top'
        remark.alignment = alignment
        visual_lines = sum(
            max(
                1,
                math.ceil(
                    sum(2 if unicodedata.east_asian_width(ch) in 'WF' else 1 for ch in line) / 66
                ),
            )
            for line in note.split('\n')
        )
        sheet.row_dimensions[row].height = min(
            300, max(sheet.row_dimensions[row].height or 0, visual_lines * 18 + 12)
        )
    Path(payload['output']).parent.mkdir(parents=True, exist_ok=True)
    book.save(payload['output'])
    return {'path': payload['output'], 'cases': len(payload['results'])}


if __name__ == '__main__':
    for stream in (sys.stdin, sys.stdout, sys.stderr):
        if hasattr(stream, 'reconfigure'):
            stream.reconfigure(encoding='utf-8')
    try:
        data = json.load(sys.stdin)
        result = {
            'write': write,
            'read': read,
            'read_editable': read_editable,
            'results': results,
        }[sys.argv[1]](data)
        # The Node parent consumes JSON, so ASCII escapes avoid Windows console codepages.
        print(json.dumps(result, ensure_ascii=True))
    except Exception as exc:
        print(f'{type(exc).__name__}: {exc}', file=sys.stderr)
        sys.exit(1)
